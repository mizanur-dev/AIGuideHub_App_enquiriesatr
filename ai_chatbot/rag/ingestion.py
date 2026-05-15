# rag_core/rag/ingestion.py

from .pdf_loader import load_pdf
from .chunking import chunk_text
from .embedding import embed_texts
from .vector_store import upsert_chunks, clear_namespace, delete_document_vectors
import threading
import logging
import fitz  # PyMuPDF
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain_core.messages import HumanMessage
import base64
from io import BytesIO
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

logger = logging.getLogger(__name__)

vision_llm = ChatGoogleGenerativeAI(
    model="gemini-2.5-flash",
    temperature=0.2,
)

def _image_to_base64(image):
    buffered = BytesIO()
    image.save(buffered, format="JPEG", quality=85)
    return base64.b64encode(buffered.getvalue()).decode("utf-8")

def describe_slide(image):
    image_b64 = _image_to_base64(image)
    
    prompt = (
        "Analyze this image from a presentation slide in detail. "
        "1. Extract ALL visible text word-for-word (OCR), especially years, titles, legislation acts, and key terms. "
        "2. If there is a chart or graph, explain the X and Y axes, legends, and data points clearly. "
        "3. If there is a table, extract and summarize the rows and columns. "
        "4. Describe what the image represents overall."
    )
    
    message = HumanMessage(
        content=[
            {"type": "text", "text": prompt},
            {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{image_b64}"}}
        ]
    )
    
    response = vision_llm.invoke([message])
    return response.content

def _extract_text_from_shape(shape, text_list, img_descriptions, slide_idx):
    if getattr(shape, "has_text_frame", False) and shape.text and shape.text.strip():
        text_list.append(shape.text.strip())

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child_shape in shape.shapes:
            _extract_text_from_shape(child_shape, text_list, img_descriptions, slide_idx)
    elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            image_stream = BytesIO(shape.image.blob)
            img = Image.open(image_stream)
            if img.mode != "RGB":
                img = img.convert("RGB")
            img_desc = describe_slide(img)
            img_descriptions.append(img_desc)
        except Exception as e:
            logger.error("Error parsing picture in slide %s: %s", slide_idx, e)


def process_pptx(file_path, namespace, original_filename):
    logger.info("Starting PPTX ingestion for %s in namespace %s", file_path, namespace)
    delete_document_vectors(namespace, original_filename)

    chunks = []
    metadatas = []

    prs = Presentation(file_path)
    for i, slide in enumerate(prs.slides):
        slide_text_list = []
        image_descriptions = []

        for shape in slide.shapes:
            _extract_text_from_shape(shape, slide_text_list, image_descriptions, i)

        slide_text = "\n".join(slide_text_list).strip()
        if not slide_text and not image_descriptions:
            logger.info("Skipping empty slide %s in %s", i + 1, original_filename)
            continue

        combined_description = f"Slide {i+1} Text:\n{slide_text}\n"
        if image_descriptions:
            combined_description += "Slide Images Descriptions:\n" + "\n".join(image_descriptions)

        chunks.append(combined_description)
        metadatas.append({
            "text": combined_description,
            "is_slide": True,
            "slide_index": i + 1,
            "file_type": "slide_deck",
            "filename": original_filename,
            "module_id": None,
            "module_name": None,
            "subsection_id": None,
            "subsection_name": None,
        })

    if not chunks:
        logger.warning("No chunks were extracted from PPTX %s", original_filename)
        return 0

    embeddings = embed_texts(chunks)
    upsert_chunks(chunks, embeddings, namespace, original_filename, custom_metadata=metadatas)
    logger.info("Successfully indexed %s PPTX slides for %s", len(chunks), original_filename)
    return len(chunks)


def process_document_async(file_path, file_type, namespace, original_filename):
    thread = threading.Thread(
        target=process_document,
        args=(file_path, file_type, namespace, original_filename),
        daemon=True
    )
    thread.start()
    return {"message": "Processing started in the background."}


def process_document(file_path, file_type, namespace, original_filename):
    is_slide_deck = file_type == "slide_deck" or original_filename.lower().endswith(('.pptx', '.ppt'))
    if is_slide_deck:
        return process_pptx(file_path, namespace, original_filename)
    return process_pdf(file_path, namespace, original_filename)


def index_structure_chunks(structure, namespace, filename):
    """Index chunks with explicit module and subsection metadata."""
    delete_document_vectors(namespace, filename)

    chunks = []
    metadatas = []

    for module in structure:
        module_id = module.get("module_id")
        module_name = module.get("module_name") or module.get("module")
        module_order = module.get("module_order", 0)

        for subsection in module.get("subsections", []):
            subsection_id = subsection.get("subsection_id")
            subsection_name = subsection.get("name")
            subsection_order = subsection.get("order", 0)
            subsection_text = subsection.get("content", "").strip()
            if not subsection_text:
                continue

            subsection_chunks = chunk_text(subsection_text)
            for chunk in subsection_chunks:
                chunks.append(chunk)
                metadatas.append({
                    "text": chunk,
                    "is_slide": False,
                    "filename": filename,
                    "module_id": module_id,
                    "module_name": module_name,
                    "module_order": module_order,
                    "subsection_id": subsection_id,
                    "subsection_name": subsection_name,
                    "subsection_order": subsection_order,
                    "file_type": "pdf",
                })

    if not chunks:
        return 0

    embeddings = embed_texts(chunks)
    upsert_chunks(chunks, embeddings, namespace=namespace, filename=filename, custom_metadata=metadatas)
    return len(chunks)


def process_pdf(file, session_id, original_filename, structure=None):
    """Process a PDF file, with optional module/subsection structure for scoped indexing."""
    if structure:
        return index_structure_chunks(structure, session_id, original_filename)

    # Legacy fallback: index the entire PDF as a single document namespace
    delete_document_vectors(session_id, original_filename)

    # 1. Load
    text = load_pdf(file)

    # 2. Chunk
    chunks = chunk_text(text)

    # 3. Embed
    embeddings = embed_texts(chunks)

    # 4. Store (namespace = session)
    upsert_chunks(chunks, embeddings, namespace=session_id, filename=original_filename)

    return len(chunks)